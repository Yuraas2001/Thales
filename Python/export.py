import sys
import mysql.connector
import os
import datetime
from reportlab.lib.pagesizes import landscape, inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font

def export_to_pdf(data, username):
    # Get the current date
    current_date = datetime.datetime.now().strftime("%d_%m_%Y")

    # Get the filename
    filename = f"{username}_{current_date}"
    counter = 1
    while os.path.exists(f"{filename} ({counter}).pdf"):
        counter += 1
    filename = f"{filename} ({counter}).pdf"

    # Calculate column widths
    col_widths = []
    for i in range(len(data[0])):
        max_length = max(len(str(row[i])) for row in data)
        col_widths.append(max_length * 6)  # 6 is an arbitrary multiplier to convert string length to points

    # Calculate page width
    page_width = sum(col_widths)

    # Create a PDF with custom page size in landscape orientation
    doc = SimpleDocTemplate(filename, pagesize=landscape((page_width, 12*inch)), topMargin=0.5*inch, bottomMargin=0.5*inch)

    # Styles for title
    styles = getSampleStyleSheet()
    title = Paragraph(f"{username} - {current_date}", styles['Title'])

    # Add column headers
    table_data = [['Programme', 'Phase', 'Description', 'Mots Clés']]
    table_data.extend(data)  # Add data rows

    # Create the table with calculated column widths
    table = Table(table_data, colWidths=col_widths)

    # Add style to the table
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.gray),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
    ]))

    # Build the PDF with title and table
    elements = [title, Spacer(1, 12), table]
    doc.build(elements)

def export_to_excel(data, username):
    # Get the current date
    current_date = datetime.datetime.now().strftime("%d_%m_%Y")

    # Get the filename
    filename = f"{username}_{current_date}"
    counter = 1
    while os.path.exists(f"{filename} ({counter}).xlsx"):
        counter += 1
    filename = f"{filename} ({counter}).xlsx"

    # Create an Excel workbook
    wb = Workbook()
    ws = wb.active

    # Add title
    title = f"{username} - {current_date}"
    ws.append([title])  # Append the title row
    ws.append([])  # Append an empty row for spacing

    # Add headers
    headers = ['Programme', 'Phase', 'Description', 'Mots Clés']
    ws.append(headers)

    # Add data to the worksheet
    for row in data:
        ws.append(row)

    # Adjust column widths
    for column in ws.columns:
        max_length = 0
        column = [cell for cell in column]
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column[0].column_letter].width = adjusted_width

    # Save the workbook
    wb.save(filename)

def main():
    # Initialize variables for parameters
    username = None
    keywords = None
    programs = None
    phase = None
    format = None

    # Parse command-line arguments
    for i in range(1, len(sys.argv), 2):
        if sys.argv[i] == '-u':
            username = sys.argv[i+1]
        elif sys.argv[i] == '-k':
            keywords = sys.argv[i+1].split(',')
        elif sys.argv[i] == '-p':
            programs = sys.argv[i+1].split(',')
        elif sys.argv[i] == '-ph':
            phase = sys.argv[i+1]
        elif sys.argv[i] == '-f':
            format = sys.argv[i+1]

    print("Username:", username)
    print("Keywords:", keywords)
    print("Programs:", programs)
    print("Phase:", phase)
    print("Format:", format)

    # Connect to the database
    try:
        conn = mysql.connector.connect(
            host="localhost",
            user="root",
            password="",
            database="BP"
        )

        # Prepare SQL query
        sql = """
            SELECT DISTINCT BonnesPratiques.IDBonnePratique, 
                            Phases.NomPhase, 
                            BonnesPratiques.Description, 
                            GROUP_CONCAT(DISTINCT Programmes.NomProgramme SEPARATOR ', ') AS Programs, 
                            GROUP_CONCAT(DISTINCT MotsCles.NomMotsCles SEPARATOR ', ') AS Keywords
            FROM BonnesPratiques
            LEFT JOIN PratiqueProg ON BonnesPratiques.IDBonnePratique = PratiqueProg.IDBonnePratique
            LEFT JOIN Programmes ON PratiqueProg.IDProgramme = Programmes.IDProgramme
            LEFT JOIN PratiquePhases ON BonnesPratiques.IDBonnePratique = PratiquePhases.IDBonnePratique
            LEFT JOIN Phases ON PratiquePhases.IDPhase = Phases.IDPhase
            LEFT JOIN PratiqueMotsCles ON BonnesPratiques.IDBonnePratique = PratiqueMotsCles.IDBonnePratique
            LEFT JOIN MotsCles ON PratiqueMotsCles.IDMotsCles = MotsCles.IDMotsCles
            WHERE 1=1
        """

        conditions = []
        params = []

        if keywords:
            placeholders = ', '.join(['%s'] * len(keywords))
            conditions.append(f"MotsCles.NomMotsCles IN ({placeholders})")
            params.extend(keywords)

        if programs:
            placeholders = ', '.join(['%s'] * len(programs))
            conditions.append(f"(Programmes.NomProgramme IN ({placeholders}) OR Programmes.NomProgramme = 'GENERIC')")
            params.extend(programs)

        if phase:
            conditions.append("Phases.NomPhase = %s")
            params.append(phase)

        if conditions:
            sql += " AND " + " AND ".join(conditions)

        sql += " GROUP BY BonnesPratiques.IDBonnePratique HAVING COUNT(DISTINCT Phases.NomPhase) = 1"

        print("SQL Query:", sql)
        print("Parameters:", params)

        # Execute the query
        cursor = conn.cursor()
        cursor.execute(sql, params)
        rows = cursor.fetchall()

        # Store the results
        results = []
        for row in rows:
            results.append([row[3], row[1], row[2], row[4]])

        # Close the database connection
        conn.close()

        # Export to the specified format
        if format == 'PDF':
            export_to_pdf(results, username)
        elif format == 'Excel':
            export_to_excel(results, username)
        else:
            print("Unsupported format:", format)

    except mysql.connector.Error as e:
        print("MySQL Error:", e)
        return

if __name__ == "__main__":
    main()
